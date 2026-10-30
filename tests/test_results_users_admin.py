import csv
import io
import json
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AttemptStatus, AuditLog, Difficulty, ExamAttempt, Role, User
from app.services import audit_admin
from app.services import exam
from app.services import results_admin as ra
from app.services import user_admin as ua
from app.services.item_import import ImportFormatError
from app.services.security import verify_password
from app.services.timeutil import KST, fmt_kst, kst_day_start_utc, parse_date
from tests.helpers import domain, make_blueprint, make_item, make_user


# =============== 시간 유틸 ===============
def test_kst_conversion_and_format():
    utc = datetime(2026, 10, 30, 15, 30, tzinfo=timezone.utc)
    assert fmt_kst(utc) == "2026-10-31 00:30" and fmt_kst(None) == "-"
    assert fmt_kst(utc, "%H:%M:%S") == "00:30:00"


def test_kst_day_boundaries_in_utc():
    start = kst_day_start_utc(date(2026, 10, 31))
    assert start == datetime(2026, 10, 30, 15, 0, tzinfo=timezone.utc)
    assert parse_date("2026-10-31") == date(2026, 10, 31)
    assert parse_date("not-a-date") is None and parse_date("") is None and parse_date(None) is None


# =============== CSV/Excel 인젝션 방어 ===============
@pytest.mark.parametrize("v", ["=1+1", "+cmd", "-2+3", "@SUM(A1)", "\tcmd", "\rcmd", "=HYPERLINK(\"http://evil\")"])
def test_safe_cell_neutralizes_formula_prefixes(v):
    assert ra.safe_cell(v) == "'" + v


@pytest.mark.parametrize("v", ["홍길동", "abc", "", 12, 3.5, None, "a=b"])
def test_safe_cell_leaves_normal_values(v):
    assert ra.safe_cell(v) == v


# =============== 결과 조회/내보내기 ===============
@pytest.fixture()
def results(db):
    d = domain(db, "ARCH")
    make_blueprint(db, {("ARCH", Difficulty.MID): (2, 50.0)}, recent_exclude_n=0)
    for _ in range(2):
        make_item(db, d, correct=("A",))
    people = {
        "good": make_user(db, "good.user"),
        "bad": make_user(db, "bad_user"),
        "evil": make_user(db, "evil"),
    }
    people["evil"].name = "=cmd|' /C calc'!A0"
    out = {}
    base = datetime(2026, 10, 30, 0, 0, tzinfo=timezone.utc)
    for i, (k, u) in enumerate(people.items()):
        a, _ = exam.start_attempt(db, u, base + timedelta(days=i))
        if k == "good":
            for ai in a.items:
                exam.save_answer(db, a, ai.seq, ["A"], base + timedelta(days=i, minutes=5))
        exam.submit_attempt(db, a, base + timedelta(days=i, minutes=10))
        out[k] = a
    db.commit()
    return db, out


def ids(rows):
    return [a.id for a, _, _ in rows]


def test_list_all_newest_first(results):
    db, a = results
    rows, total = ra.list_attempts(db, ra.Filters())
    assert total == 3 and ids(rows) == [a["evil"].id, a["bad"].id, a["good"].id]


def test_filter_by_passed_status_and_query(results):
    db, a = results
    assert ids(ra.list_attempts(db, ra.Filters(passed="1"))[0]) == [a["good"].id]
    assert set(ids(ra.list_attempts(db, ra.Filters(passed="0"))[0])) == {a["bad"].id, a["evil"].id}
    assert ra.list_attempts(db, ra.Filters(status="IN_PROGRESS"))[1] == 0
    assert ids(ra.list_attempts(db, ra.Filters(q="GOOD"))[0]) == [a["good"].id]  # 대소문자 무시
    assert ra.list_attempts(db, ra.Filters(q="%"))[1] == 0 and ra.list_attempts(db, ra.Filters(q="_"))[1] == 1  # 'bad_user' 의 밑줄만 일치
    assert ra.list_attempts(db, ra.Filters(status="BOGUS"))[1] == 3  # 알 수 없는 값은 무시


def test_filter_by_kst_date_range(results):
    db, a = results
    # 시작 시각(UTC): 10/30 00:00, 10/31 00:00, 11/01 00:00 = KST 09:00 같은 날들
    only_31 = ra.Filters(date_from=date(2026, 10, 31), date_to=date(2026, 10, 31))
    assert ids(ra.list_attempts(db, only_31)[0]) == [a["bad"].id]
    assert ra.list_attempts(db, ra.Filters(date_from=date(2026, 11, 2)))[1] == 0
    assert ra.list_attempts(db, ra.Filters(date_to=date(2026, 10, 29)))[1] == 0


def test_kst_date_filter_uses_kst_not_utc_day(db):
    d = domain(db, "ARCH")
    make_blueprint(db, {("ARCH", Difficulty.MID): (1, 100.0)})
    make_item(db, d)
    u = make_user(db)
    # UTC 10/30 16:00 = KST 10/31 01:00 → KST 기준 10/31 에 속한다
    a, _ = exam.start_attempt(db, u, datetime(2026, 10, 30, 16, 0, tzinfo=timezone.utc))
    assert ra.list_attempts(db, ra.Filters(date_from=date(2026, 10, 31), date_to=date(2026, 10, 31)))[1] == 1
    assert ra.list_attempts(db, ra.Filters(date_from=date(2026, 10, 30), date_to=date(2026, 10, 30)))[1] == 0
    assert a.id


def test_pagination(results):
    db, _ = results
    rows, total = ra.list_attempts(db, ra.Filters(), page=2, per_page=2)
    assert total == 3 and len(rows) == 1


def test_summary(results):
    db, _ = results
    s = ra.summary(db, ra.Filters())
    assert s["n"] == 3 and s["passed"] == 1 and s["pass_rate"] == pytest.approx(100 / 3) and s["avg"] == pytest.approx(100 / 3)
    empty = ra.summary(db, ra.Filters(q="nobody"))
    assert empty == {"n": 0, "passed": 0, "pass_rate": None, "avg": None}


def test_export_table_content_and_order(results):
    db, a = results
    header, rows = ra.export_table(db, ra.Filters())
    assert header[:10] == ["회차", "로그인 ID", "이름", "Blueprint", "시작(KST)", "제출(KST)", "소요(분)", "상태", "총점", "합격"]
    assert any("ARCH" in h for h in header[10:])
    assert [r[0] for r in rows] == sorted(r[0] for r in rows)
    good = next(r for r in rows if r[1] == "good.user")
    assert good[7] == "제출" and good[8] == 100 and good[9] == "합격" and good[6] == 10.0
    assert good[4] == "2026-10-30 09:00:00"  # UTC 00:00 → KST 09:00


def test_csv_export_has_bom_and_is_formula_safe(results):
    db, _ = results
    header, rows = ra.export_table(db, ra.Filters())
    raw = ra.to_csv(header, rows)
    assert raw.startswith(b"\xef\xbb\xbf")
    parsed = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    names = [r[2] for r in parsed[1:]]
    assert "'=cmd|' /C calc'!A0" in names and "=cmd|' /C calc'!A0" not in names
    assert all(not c.startswith(ra.FORMULA_PREFIXES) for r in parsed for c in r if isinstance(c, str))


def test_xlsx_export_is_formula_safe_and_loadable(results):
    db, _ = results
    header, rows = ra.export_table(db, ra.Filters())
    wb = load_workbook(io.BytesIO(ra.to_xlsx(header, rows)))
    ws = wb.active
    assert ws.title == "응시 결과" and ws.freeze_panes == "A2"
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any(c.data_type == "f" for c in cells)  # 수식으로 저장된 셀이 없어야 한다
    assert "'=cmd|' /C calc'!A0" in [c.value for c in ws["C"]]
    assert ws.max_row == 4 and ws["A1"].value == "회차"


# =============== 응시자 계정 ===============
@pytest.fixture()
def admin(db):
    return make_user(db, role=Role.ADMIN)


def test_generate_password_properties():
    pws = {ua.generate_password() for _ in range(300)}
    assert len(pws) == 300  # 충돌 없음
    for pw in pws:
        assert len(pw) == 12
        assert any(c.isupper() for c in pw) and any(c.islower() for c in pw) and any(c.isdigit() for c in pw)
        assert not set(pw) & set("0O1lI")


def test_reset_password_changes_hash_clears_lock_and_logs(db, admin):
    u = make_user(db, "locked.user")
    old = u.password_hash
    u.failed_login_count, u.locked_until = 3, datetime.now(timezone.utc) + timedelta(minutes=5)
    cred = ua.reset_password(db, u, admin)
    assert u.password_hash != old and verify_password(cred.password, u.password_hash)
    assert u.failed_login_count == 0 and u.locked_until is None
    assert db.scalar(select(AuditLog).where(AuditLog.action == "USER_PASSWORD_RESET")).target == f"user:{u.id}"


def test_unlock(db, admin):
    u = make_user(db)
    u.failed_login_count, u.locked_until = 0, datetime.now(timezone.utc) + timedelta(minutes=5)
    ua.unlock(db, u, admin)
    assert u.locked_until is None and db.scalar(select(AuditLog).where(AuditLog.action == "USER_UNLOCK"))


def test_cannot_manage_admins_or_self(db, admin):
    other = make_user(db, role=Role.ADMIN)
    for fn in (ua.reset_password, ua.unlock):
        with pytest.raises(ua.UserAdminError, match="관리자 계정"):
            fn(db, other, admin)
        with pytest.raises(ua.UserAdminError):
            fn(db, admin, admin)


def test_list_users_filters(db, admin):
    make_user(db, "alpha.one")
    make_user(db, "beta.two")
    assert ua.list_users(db)[1] == 3
    assert [u.login_id for u in ua.list_users(db, q="ALPHA")[0]] == ["alpha.one"]
    assert ua.list_users(db, role="ADMIN")[1] == 1 and ua.list_users(db, q="%")[1] == 0


# =============== 감사 로그 조회 ===============
def test_audit_listing_filters(db, admin):
    from app.services.audit import log_audit

    log_audit(db, admin.id, "ITEM_APPROVE", "item:1", {"a": 1})
    log_audit(db, None, "LOGIN_FAILED", "ghost", None)
    log_audit(db, admin.id, "BLUEPRINT_UPDATE", "blueprint:2", None)
    db.flush()
    rows, total = audit_admin.list_audit(db)
    assert total == 3 and [r[0].action for r in rows][0] == "BLUEPRINT_UPDATE"
    assert audit_admin.list_audit(db, action="LOGIN_FAILED")[1] == 1
    assert audit_admin.list_audit(db, q="item:1")[1] == 1
    login = admin.login_id
    assert audit_admin.list_audit(db, q=login)[1] == 2  # 사용자 로그인 ID 로도 검색
    assert audit_admin.list_audit(db, q="%")[1] == 0
    assert audit_admin.distinct_actions(db) == ["BLUEPRINT_UPDATE", "ITEM_APPROVE", "LOGIN_FAILED"]
    today = datetime.now(KST).date()
    assert audit_admin.list_audit(db, date_from=today, date_to=today)[1] == 3
    assert audit_admin.list_audit(db, date_to=today - timedelta(days=1))[1] == 0
    login_failed = audit_admin.list_audit(db, action="LOGIN_FAILED")[0][0]
    assert login_failed[1] is None  # 사용자 없는 기록은 login 이 None


