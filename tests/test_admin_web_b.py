import io
import json
import re

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    AttemptStatus, AuditLog, Blueprint, Difficulty, Domain, ExamAttempt, Item, ItemOverride, ItemStatus, User,
)
from tests.test_web import HX, PW, client, env, login, page_token, token  # noqa: F401


@pytest.fixture()
def Session(env):  # noqa: F811
    return env[1]


@pytest.fixture()
def adm(client):  # noqa: F811
    login(client, "admin")
    return client


def tok(c):
    return token(c.get("/admin").text)


def post(c, path, data=None, **kw):
    return c.post(path, data={"csrf_token": tok(c), **(data or {})}, **kw)


def audit_actions(Session):
    with Session() as db:
        return [a.action for a in db.scalars(select(AuditLog).order_by(AuditLog.id))]


def finished_attempt(Session, login_id="cand", answer=True):
    from app.services import exam

    with Session() as db:
        u = db.scalar(select(User).where(User.login_id == login_id))
        a, _ = exam.start_attempt(db, u)
        if answer:
            for ai in a.items:
                exam.save_answer(db, a, ai.seq, list(ai.item.correct_keys))
        exam.submit_attempt(db, a)
        return a.id


# ================= 접근 제어 / CSRF =================
B_GET = ["/admin/blueprints", "/admin/blueprints/new", "/admin/blueprints/1", "/admin/results", "/admin/results/1",
         "/admin/results/export.csv", "/admin/results/export.xlsx", "/admin/users", "/admin/settings", "/admin/audit"]
B_POST = ["/admin/blueprints/new", "/admin/blueprints/1", "/admin/blueprints/1/activate",
          "/admin/users/2/approve", "/admin/users/2/reject", "/admin/users/2/reset-password", "/admin/users/2/unlock", "/admin/settings",
          "/admin/items/1/override", "/admin/items/1/override/revoke"]


def test_anonymous_redirected(client):  # noqa: F811
    for p in B_GET:
        r = client.get(p, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/login", p


def test_candidate_forbidden(client):  # noqa: F811
    login(client, "cand")
    for p in B_GET:
        assert client.get(p).status_code == 403, p
    t = page_token(client)
    for p in B_POST:
        assert client.post(p, data={"csrf_token": t}).status_code == 403, p


def test_posts_require_csrf(adm):
    for p in B_POST:
        assert adm.post(p).status_code == 403, p
        assert adm.post(p, data={"csrf_token": "forged"}).status_code == 403, p


# ================= Blueprint =================
def bp_form(**over):
    d = {"name": "웹 기준표", "time_limit_min": "30", "pass_score": "70", "domain_min_percent": "40",
         "recent_exclude_n": "1", "sort_by_difficulty": "on",
         "count_ARCH_LOW": "10", "points_ARCH_LOW": "2", "count_ARCH_MID": "16", "points_ARCH_MID": "2.5",
         "count_OPS_HIGH": "8", "points_OPS_HIGH": "5"}
    d.update(over)
    return d


def test_blueprint_list_and_detail_matrix(adm):
    html = adm.get("/admin/blueprints").text
    assert "기본 40문항" in html and "활성" in html and "충분" in html
    page = adm.get("/admin/blueprints/1").text
    assert page.count('name="count_') == 21 and page.count('name="points_') == 21
    assert "40문항" in page and "100점" in page and 'class="short"' not in page


def test_blueprint_matrix_shows_red_cells_for_shortage(adm, Session):
    with Session() as db:
        arch = db.scalar(select(Domain).where(Domain.code == "ARCH"))
        for it in db.scalars(select(Item).where(Item.domain_id == arch.id, Item.difficulty == Difficulty.HIGH)):
            it.status = ItemStatus.RETIRED
        db.commit()
    page = adm.get("/admin/blueprints/1").text
    assert 'class="short"' in page and "<b>부족</b>" in page
    assert "부족 1개 셀" in adm.get("/admin/blueprints").text


def test_blueprint_update_valid_and_audited(adm, Session):
    form = {"name": "기본 40문항", "time_limit_min": "75", "pass_score": "70", "domain_min_percent": "40",
            "recent_exclude_n": "2", "sort_by_difficulty": "on"}
    with Session() as db:
        bp = db.get(Blueprint, 1)
        for r in bp.rules:
            form[f"count_{r.domain.code}_{r.difficulty.value}"] = str(r.item_count)
            form[f"points_{r.domain.code}_{r.difficulty.value}"] = str(r.points_per_item)
    r = post(adm, "/admin/blueprints/1", form, follow_redirects=True)
    assert "저장했습니다" in r.text and "이후 시작하는 시험부터 적용" in r.text
    with Session() as db:
        assert db.get(Blueprint, 1).time_limit_min == 75
    assert "BLUEPRINT_UPDATE" in audit_actions(Session)


def test_blueprint_update_rejects_bad_sum_and_keeps_input(adm, Session):
    form = bp_form(name="기본 40문항", points_ARCH_LOW="3")
    r = post(adm, "/admin/blueprints/1", form)
    assert r.status_code == 422 and "배점 합계가" in r.text and 'value="3"' in r.text
    with Session() as db:
        assert db.get(Blueprint, 1).time_limit_min == 60  # 저장되지 않음


def test_blueprint_new_prefilled_and_create_then_activate(adm, Session):
    page = adm.get("/admin/blueprints/new").text
    assert 'name="name" value=""' in page  # 이름은 비움
    assert 'name="count_EXTRACT_MID" value="3"' in page and 'name="count_REPLICAT_MID" value="4"' in page  # 활성 규칙 복사
    # 가용 범위(셀당 3개) 안에서 합계 100점: ARCH LOW 3×10 + MID 3×10 + HIGH 2×20
    ok = bp_form(count_ARCH_LOW="3", points_ARCH_LOW="10", count_ARCH_MID="3", points_ARCH_MID="10",
                 count_OPS_HIGH="0", points_OPS_HIGH="0", count_ARCH_HIGH="2", points_ARCH_HIGH="20")
    r = post(adm, "/admin/blueprints/new", ok, follow_redirects=False)
    assert r.status_code == 303
    new_id = int(r.headers["location"].rsplit("/", 1)[1])
    with Session() as db:
        bp = db.get(Blueprint, new_id)
        assert bp.total_items == 8 and not bp.is_active and bp.time_limit_min == 30
    r = post(adm, f"/admin/blueprints/{new_id}/activate", follow_redirects=True)
    assert "활성화했습니다" in r.text
    with Session() as db:
        assert [b.id for b in db.scalars(select(Blueprint).where(Blueprint.is_active.is_(True)))] == [new_id]


def test_blueprint_activation_blocked_on_shortage_with_message(adm, Session):
    r = post(adm, "/admin/blueprints/new", bp_form(count_ARCH_LOW="10", points_ARCH_LOW="2"), follow_redirects=False)
    new_id = int(r.headers["location"].rsplit("/", 1)[1])  # ARCH/LOW 10문항 필요 — 가용 3
    page = adm.get(f"/admin/blueprints/{new_id}").text
    assert "활성화할 수 없습니다" in page and "disabled" in page
    r = post(adm, f"/admin/blueprints/{new_id}/activate", follow_redirects=True)
    assert "ARCH/LOW" in r.text and "활성화할 수 없습니다" in r.text
    with Session() as db:
        assert db.get(Blueprint, 1).is_active and not db.get(Blueprint, new_id).is_active


def test_blueprint_duplicate_name_and_unknown_id(adm):
    r = post(adm, "/admin/blueprints/new", bp_form(name="기본 40문항"))
    assert r.status_code == 422 and "같은 이름" in r.text
    assert adm.get("/admin/blueprints/999").status_code == 404
    assert post(adm, "/admin/blueprints/999/activate").status_code == 404


def test_active_blueprint_change_applies_to_next_exam(adm, client, Session):  # noqa: F811
    # 제한시간 변경 → 이후 시작한 시험의 만료 시각에 반영
    form = {"name": "기본 40문항", "time_limit_min": "10", "pass_score": "70", "domain_min_percent": "40", "recent_exclude_n": "2"}
    with Session() as db:
        for r in db.get(Blueprint, 1).rules:
            form[f"count_{r.domain.code}_{r.difficulty.value}"] = str(r.item_count)
            form[f"points_{r.domain.code}_{r.difficulty.value}"] = str(r.points_per_item)
    post(adm, "/admin/blueprints/1", form)
    from app.services import exam

    with Session() as db:
        u = db.scalar(select(User).where(User.login_id == "cand"))
        a, _ = exam.start_attempt(db, u)
        assert round((a.expires_at - a.started_at).total_seconds() / 60) == 10


# ================= 결과 / 내보내기 =================
def test_results_list_filters_and_summary(adm, Session):
    a1 = finished_attempt(Session, "cand", answer=True)
    a2 = finished_attempt(Session, "cand2", answer=False)
    html = adm.get("/admin/results").text
    assert "(2건)" in html and "100.0" in html and "50.0%" in html and f"#{a1}" in html
    assert "(1건)" in adm.get("/admin/results?passed=1").text and f"/admin/results/{a1}" in adm.get("/admin/results?passed=1").text
    assert "(1건)" in adm.get("/admin/results?q=cand2").text
    assert "(0건)" in adm.get("/admin/results?status=IN_PROGRESS").text
    assert adm.get("/admin/results?date_from=garbage&page=-3").status_code == 200
    assert a2


def test_result_detail_shows_items_and_cutoff(adm, Session):
    aid = finished_attempt(Session, "cand", answer=False)
    page = adm.get(f"/admin/results/{aid}").text
    assert "불합격" in page and "과락" in page and page.count("미응답") == 40
    assert "KST" in page and "seed" in page
    assert adm.get("/admin/results/99999").status_code == 404


def test_export_csv_bom_headers_and_audit(adm, Session):
    finished_attempt(Session, "cand", answer=True)
    r = adm.get("/admin/results/export.csv?passed=1")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert re.match(r'attachment; filename="results_\d{8}_\d{4}\.csv"', r.headers["content-disposition"])
    assert r.content.startswith(b"\xef\xbb\xbf") and "cand" in r.content.decode("utf-8-sig") and "합격" in r.content.decode("utf-8-sig")
    with Session() as db:
        log = db.scalars(select(AuditLog).where(AuditLog.action == "RESULT_EXPORT")).one()
        assert log.target == "csv" and log.detail["rows"] == 1 and "passed=1" in log.detail["filters"]


def test_export_xlsx_loadable_and_filtered(adm, Session):
    finished_attempt(Session, "cand", answer=True)
    finished_attempt(Session, "cand2", answer=False)
    r = adm.get("/admin/results/export.xlsx?q=cand2")
    assert "spreadsheetml" in r.headers["content-type"] and ".xlsx" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 2 and ws["B2"].value == "cand2"


def test_export_neutralizes_formula_in_user_name(adm, Session):
    with Session() as db:
        db.scalar(select(User).where(User.login_id == "cand")).name = "=HYPERLINK(\"http://evil\",\"x\")"
        db.commit()
    finished_attempt(Session, "cand", answer=True)
    csv_text = adm.get("/admin/results/export.csv").content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and ",=HYPERLINK" not in csv_text
    ws = load_workbook(io.BytesIO(adm.get("/admin/results/export.xlsx").content)).active
    assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)


def test_export_links_carry_filters(adm, Session):
    finished_attempt(Session, "cand")
    html = adm.get("/admin/results?passed=1&q=cand").text
    assert "/admin/results/export.csv?passed=1&amp;q=cand" in html or "/admin/results/export.csv?passed=1&q=cand" in html


def test_dates_are_shown_in_kst(adm, Session, client):  # noqa: F811
    from datetime import timedelta

    aid = finished_attempt(Session, "cand")
    with Session() as db:
        a = db.get(ExamAttempt, aid)
        expected = (a.started_at + timedelta(hours=9)).strftime("%Y-%m-%d %H:%M")
    assert expected in adm.get("/admin/results").text
    client.cookies.clear()
    login(client, "cand")
    assert expected in client.get("/").text and "UTC" not in client.get("/").text


# ================= 응시자 계정 =================


def test_reset_password_flow_old_password_stops_working(adm, client, Session):  # noqa: F811
    with Session() as db:
        uid = db.scalar(select(User).where(User.login_id == "cand")).id
    r = post(adm, f"/admin/users/{uid}/reset-password")
    new_pw = re.search(r'class="secret">([^<]+)<', r.text).group(1)
    assert "비밀번호 재설정" in r.text and r.headers["cache-control"] == "no-store"
    client.cookies.clear()
    assert login(client, "cand", PW).status_code == 401
    assert login(client, "cand", new_pw).status_code == 303


def test_unlock_locked_account(adm, client, Session):  # noqa: F811
    client.cookies.clear()
    for _ in range(5):
        login(client, "cand2", "wrong-password")
    assert login(client, "cand2", PW).status_code == 401
    client.cookies.clear()
    login(client, "admin")
    users = adm.get("/admin/users").text
    assert "잠김" in users and "잠금 해제" in users
    with Session() as db:
        uid = db.scalar(select(User).where(User.login_id == "cand2")).id
    assert "잠금을 해제했습니다" in post(adm, f"/admin/users/{uid}/unlock", follow_redirects=True).text
    client.cookies.clear()
    assert login(client, "cand2", PW).status_code == 303


def test_cannot_reset_admin_account_via_ui(adm, Session):
    with Session() as db:
        admin_id = db.scalar(select(User).where(User.login_id == "admin")).id
    r = post(adm, f"/admin/users/{admin_id}/reset-password", follow_redirects=True)
    assert "관리자 계정" in r.text or "본인 계정" in r.text
    assert "secret" not in r.text
    assert post(adm, "/admin/users/99999/reset-password").status_code == 404


def test_user_list_filters_and_hides_admin_actions(adm):
    html = adm.get("/admin/users").text
    assert "cand2" in html and "관리자" in html
    only_admin = adm.get("/admin/users?role=ADMIN").text
    assert "cand2" not in only_admin and "비밀번호 재설정" not in only_admin


# ================= 설정 =================
def test_settings_toggle_explanations_affects_candidate_result(adm, client, Session):  # noqa: F811
    aid = finished_attempt(Session, "cand", answer=False)
    page = adm.get("/admin/settings").text
    assert re.search(r'name="show_explanations"\s+checked', page)
    assert not re.search(r'name="require_separate_approver"\s+checked', page)
    r = post(adm, "/admin/settings", {}, follow_redirects=True)  # 체크박스 모두 해제 = 둘 다 꺼짐
    assert "설정을 저장했습니다" in r.text
    with Session() as db:
        entry = db.scalars(select(AuditLog).where(AuditLog.action == "SETTINGS_UPDATE")).one()
        assert entry.detail["changed"] == {"show_explanations": [True, False]}
    client.cookies.clear()
    login(client, "cand")
    assert "공개되지 않습니다" in client.get(f"/exam/{aid}/result").text
    client.cookies.clear()
    login(client, "admin")
    post(adm, "/admin/settings", {"show_explanations": "on"})
    client.cookies.clear()
    login(client, "cand")
    assert 'class="explain"' in client.get(f"/exam/{aid}/result").text


def test_settings_no_change_message_and_no_audit(adm, Session):
    post(adm, "/admin/settings", {"show_explanations": "on"})  # 기본값과 동일
    assert "변경된 내용이 없습니다" in post(adm, "/admin/settings", {"show_explanations": "on"}, follow_redirects=True).text
    assert "SETTINGS_UPDATE" not in audit_actions(Session)


def test_four_eyes_setting_blocks_self_approval_in_ui(adm, Session):
    post(adm, "/admin/settings", {"show_explanations": "on", "require_separate_approver": "on"})
    data = {"domain": "OPS", "difficulty": "LOW", "item_type": "SINGLE", "architecture": "BOTH", "ogg_version": "19c",
            "stem": "분리 승인 테스트", "option_A": "가", "option_B": "나", "correct_keys": ["A"], "explanation": "해설",
            "variant_group": "FE-1"}
    r = post(adm, "/admin/items/new", data, follow_redirects=False)
    iid = int(r.headers["location"].rsplit("/", 1)[1])
    post(adm, f"/admin/items/{iid}/submit-review")
    r = post(adm, f"/admin/items/{iid}/approve", follow_redirects=True)
    assert "작성자 본인은 승인할 수 없습니다" in r.text
    with Session() as db:
        assert db.get(Item, iid).status == ItemStatus.REVIEWED


# ================= 감사 로그 =================
def test_audit_page_lists_and_filters(adm, client, Session):  # noqa: F811
    login_failed = login  # noqa: F841
    client.cookies.clear()
    login(client, "cand", "wrong-password")
    login(client, "admin")
    page = adm.get("/admin/audit").text
    assert "LOGIN_FAILED" in page and "LOGIN_OK" in page and "감사 로그" in page
    only = adm.get("/admin/audit?action=LOGIN_FAILED").text
    assert "LOGIN_FAILED" in only and "LOGIN_OK</td>" not in only
    assert "(0건)" in adm.get("/admin/audit?q=zzz-none").text
    assert adm.get("/admin/audit?date_from=bad&page=0").status_code == 200


def test_audit_page_escapes_attacker_controlled_values(adm, client):  # noqa: F811
    client.cookies.clear()
    login(client, '<script>alert(1)</script>', "x")  # 로그인 ID 가 target 으로 기록됨
    login(client, "admin")
    html = adm.get("/admin/audit").text
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html


def test_audit_detail_rendered_as_readable_korean_json(adm, Session):
    aid = finished_attempt(Session, "cand", answer=False)
    post(adm, "/admin/items/1/override", {"mode": "CREDIT_ALL", "reason": "정답 키 오류 확인"})
    html = adm.get("/admin/audit?action=RESULT_OVERRIDE_SET").text
    assert "정답 키 오류 확인" in html and "\\u" not in html and "regraded" in html
    assert aid


# ================= 채점 보정(결과 수정) =================
def test_override_via_ui_regrades_results_and_notifies_candidate(adm, client, Session):  # noqa: F811
    aid = finished_attempt(Session, "cand", answer=False)  # 0점
    with Session() as db:
        a = db.get(ExamAttempt, aid)
        target = a.items[0].item_id
        pts = a.items[0].points
    page = adm.get(f"/admin/items/{target}").text
    assert "채점 보정 (결과 수정)" in page
    r = post(adm, f"/admin/items/{target}/override", {"mode": "CREDIT_ALL", "reason": "정답 키 오류"}, follow_redirects=True)
    assert "채점 보정을 적용했습니다: 재채점 1건 · 점수 변경 1건" in r.text and "정답 인정" in r.text
    with Session() as db:
        a = db.get(ExamAttempt, aid)
        assert a.total_score == pytest.approx(pts) and a.status == AttemptStatus.SUBMITTED
        assert db.get(ItemOverride, target).mode.value == "CREDIT_ALL"
    detail = adm.get(f"/admin/results/{aid}").text
    assert "정답 인정(보정)" in detail
    client.cookies.clear()
    login(client, "cand")
    assert "관리자 채점 보정" in client.get(f"/exam/{aid}/result").text
    client.cookies.clear()
    login(client, "admin")
    r = post(adm, f"/admin/items/{target}/override/revoke", {"reason": "철회"}, follow_redirects=True)
    assert "채점 보정을 해제했습니다" in r.text
    with Session() as db:
        assert db.get(ExamAttempt, aid).total_score == 0 and db.get(ItemOverride, target) is None


def test_override_requires_reason_and_valid_mode(adm, Session):
    aid = finished_attempt(Session, "cand", answer=False)
    with Session() as db:
        target = db.get(ExamAttempt, aid).items[0].item_id
    assert "사유를 입력" in post(adm, f"/admin/items/{target}/override", {"mode": "CREDIT_ALL", "reason": " "}, follow_redirects=True).text
    assert "유형이 올바르지 않습니다" in post(adm, f"/admin/items/{target}/override", {"mode": "ALL", "reason": "x"}, follow_redirects=True).text
    assert "적용 중인 보정이 없습니다" in post(adm, f"/admin/items/{target}/override/revoke", {"reason": "x"}, follow_redirects=True).text
    with Session() as db:
        assert db.get(ItemOverride, target) is None and db.get(ExamAttempt, aid).total_score == 0


def test_override_section_hidden_for_items_without_history(adm):
    assert "채점 보정 (결과 수정)" not in adm.get("/admin/items/64").text  # 출제된 적 없는 문항
    assert adm.post("/admin/items/99999/override", data={"csrf_token": tok(adm), "mode": "EXCLUDE", "reason": "x"}).status_code == 404


def test_new_nav_links_and_dashboard_buttons(adm):
    html = adm.get("/admin").text
    for href in ("/admin/blueprints", "/admin/results", "/admin/users", "/admin/settings", "/admin/audit"):
        assert f'href="{href}"' in html
